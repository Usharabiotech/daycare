import sys, json
from collections import Counter
import openpyxl

src = sys.argv[1]
out = sys.argv[2]

wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = ws.iter_rows(values_only=True)
header = next(rows)
idx = {h: i for i, h in enumerate(header)}

def g(r, name):
    v = r[idx[name]]
    if v is None:
        return None
    s = str(v).strip()
    return s or None

status_counts = Counter()
progtype_counts = Counter()
total = 0
active = []
seen_lic = set()

for r in rows:
    total += 1
    status = g(r, 'Licence Status')
    status_counts[status] += 1
    if status != 'Active':
        continue
    lic = g(r, 'Licence Number')
    if lic and lic in seen_lic:
        continue
    if lic:
        seen_lic.add(lic)
    progtype_counts[g(r, 'Program Type Desc')] += 1
    num = g(r, 'Street Number'); st = g(r, 'Street Name'); typ = g(r, 'Street Type')
    street = ' '.join([p for p in [num, st, typ] if p])
    active.append({
        'licensee': g(r, 'Licensee Name'),
        'site': g(r, 'Child Care Site Name'),
        'programType': g(r, 'Program Type Desc'),
        'programOption': g(r, 'Program Option'),
        'licenceNo': lic,
        'issueDate': (g(r, 'Original Issue Date') or '')[:10] or None,
        'language': g(r, 'Language of Service Desc'),
        'region': g(r, 'Region Display Name'),
        'cmsm': g(r, 'CMSM / DSSAB Name'),
        'street': street or None,
        'city': g(r, 'City'),
        'province': g(r, 'Province'),
        'postal': g(r, 'Postal Code'),
    })

with open(out, 'w', encoding='utf-8') as f:
    json.dump(active, f, ensure_ascii=False)

cities = Counter(a['city'] for a in active)
licensees = set(a['licensee'] for a in active if a['licensee'])
print("Total rows:", total)
print("Status breakdown:", dict(status_counts))
print("Active facilities (deduped by licence):", len(active))
print("Unique active licensees (orgs):", len(licensees))
print("Program types:", dict(progtype_counts))
print("Distinct cities:", len(cities))
print("Top 12 cities:", cities.most_common(12))
missing_site = sum(1 for a in active if not a['site'])
missing_city = sum(1 for a in active if not a['city'])
print("Active missing site name:", missing_site, "| missing city:", missing_city)
